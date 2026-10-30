"""Synthetic Office/OCR output gates; run with requirements-test.lock and real engines."""
import io
import json
import os
from pathlib import Path
import re
import shutil
import tempfile
import time
import unittest

from docx import Document
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
import pypdfium2 as pdfium

from engines import office_to_pdf, ocr_pdf, inspect_office, verify_output


def text_of(path):
    with pdfium.PdfDocument(path) as pdf:
        result = []
        for page in pdf:
            text = page.get_textpage()
            try:
                result.append(text.get_text_range())
            finally:
                text.close()
            page.close()
        return '\n'.join(result)


def pixels_of(path):
    with pdfium.PdfDocument(path) as pdf:
        page = pdf[0]
        bitmap = page.render(scale=0.5)
        pixels = bitmap.to_pil().tobytes()
        bitmap.close()
        page.close()
        return pixels


def normalize(value):
    return re.sub(r'\s+', ' ', value).strip()


def distance(left, right):
    row = list(range(len(right)+1))
    for index, a in enumerate(left):
        next_row = [index+1]
        for other, b in enumerate(right):
            next_row.append(min(row[other+1]+1, next_row[other]+1, row[other]+(a != b)))
        row = next_row
    return row[-1]


class Acceptance(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.metrics = []
        cls.artifacts = Path(os.environ.get('DOCFLOW_ACCEPTANCE_DIR','../artifacts/processing')).resolve()
        cls.artifacts.mkdir(parents=True,exist_ok=True)

    @classmethod
    def tearDownClass(cls):
        (cls.artifacts/'acceptance.json').write_text(json.dumps(cls.metrics,indent=2,ensure_ascii=False)+'\n')

    def test_office_docx_pptx_xlsx_real_pages_tables_charts_and_text(self):
        self.assertIsNotNone(shutil.which('libreoffice'),'Acceptance requires the real LibreOffice binary')
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            doc = Document()
            doc.add_heading('DocFlow Office acceptance',0)
            doc.add_paragraph('A table, an image and a deliberate page break.')
            table = doc.add_table(rows=3,cols=2)
            for row,values in zip(table.rows,[('Item','Amount'),('Alpha','125.50'),('Beta','99.00')]):
                for cell,value in zip(row.cells,values): cell.text=value
            image_path = root/'fixture.png'
            Image.new('RGB',(300,150),(20,110,180)).save(image_path)
            doc.add_picture(str(image_path))
            doc.add_page_break()
            doc.add_heading('Second page preserved',1)
            doc.save(root/'input.docx')

            deck = Presentation()
            first = deck.slides.add_slide(deck.slide_layouts[1])
            first.shapes.title.text = 'DocFlow slide one'
            first.placeholders[1].text = 'Tables and charts stay visible'
            second = deck.slides.add_slide(deck.slide_layouts[5])
            second.shapes.title.text = 'DocFlow slide two'
            chart = CategoryChartData()
            chart.categories = ['Alpha','Beta','Gamma']
            chart.add_series('Revenue',[10,20,30])
            second.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(7), Inches(4),chart)
            deck.save(root/'input.pptx')

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = 'Revenue'
            for row in [('Category','Amount'),('Alpha',10),('Beta',20),('Gamma',30)]: sheet.append(row)
            bar = BarChart()
            bar.title = 'Revenue chart'
            bar.add_data(Reference(sheet,min_col=2,min_row=1,max_row=4),titles_from_data=True)
            bar.set_categories(Reference(sheet,min_col=1,min_row=2,max_row=4))
            sheet.add_chart(bar,'D2')
            workbook.create_sheet('Summary')['A1'] = 'DocFlow workbook summary'
            for tab in workbook:
                tab.print_area = 'A1:M24'
                tab.sheet_properties.pageSetUpPr.fitToPage=True
                tab.page_setup.fitToWidth=1
                tab.page_setup.fitToHeight=1
            workbook.save(root/'input.xlsx')

            for extension, expected in [('docx',['Second page preserved','Alpha','125.50']),('pptx',['DocFlow slide one','DocFlow slide two','Gamma']),('xlsx',['Alpha','Revenue chart','DocFlow workbook summary'])]:
                with self.subTest(extension=extension):
                    job = root/extension
                    job.mkdir()
                    (job/'tmp').mkdir()
                    source = job/f'input.{extension}'
                    shutil.copyfile(root/f'input.{extension}',source)
                    inspect_office(source,extension)
                    started = time.monotonic()
                    output = job/'result.pdf'
                    office_to_pdf(source,output,extension)
                    measured = verify_output(output,expected_pages=2)
                    text = text_of(output)
                    for phrase in expected: self.assertIn(phrase,text)
                    shutil.copyfile(output,self.artifacts/f'office-{extension}.pdf')
                    self.metrics.append({'operation':'office-to-pdf','format':extension,'pages':measured['pages'],'size':measured['size'],'processingMs':round((time.monotonic()-started)*1000)})

    def test_ocr_english_german_searchable_layer_and_original_visuals(self):
        cases = [
            ('eng',['eng'],'DocFlow makes scanned documents searchable.\nClear pages preserve every important detail.\nInvoice number 2026 1042 and total 125.50.'),
            ('deu',['eng','deu'],'DocFlow macht gescannte Dokumente durchsuchbar.\nKlare Seiten bewahren wichtige Informationen.\nRechnung Nummer 2026 1042 und Summe 125,50.'),
        ]
        font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',42)
        for name,languages,expected in cases:
            with self.subTest(language=name),tempfile.TemporaryDirectory() as temporary:
                job = Path(temporary)
                (job/'tmp').mkdir()
                image = Image.new('RGB',(1800,700),'white')
                ImageDraw.Draw(image).multiline_text((80,100),expected,font=font,fill='black',spacing=35)
                source, output = job/'input.pdf',job/'result.pdf'
                image.save(source,'PDF',resolution=200,quality=98)
                started=time.monotonic()
                ocr_pdf(source,output,{'languages':languages,'deskew':False,'rotate':False})
                actual, wanted = normalize(text_of(output)),normalize(expected)
                cer = distance(wanted,actual)/len(wanted)
                wer = distance(wanted.split(),actual.split())/len(wanted.split())
                self.assertLessEqual(cer,0.03,actual)
                self.assertLessEqual(wer,0.10,actual)
                self.assertEqual(pixels_of(source),pixels_of(output),'Adding a layer changed the visible scan')
                shutil.copyfile(output,self.artifacts/f'ocr-{name}.pdf')
                self.metrics.append({'operation':'ocr','language':name,'cer':cer,'wer':wer,'visualPixelsIdentical':True,'processingMs':round((time.monotonic()-started)*1000)})


if __name__ == '__main__':
    unittest.main()
